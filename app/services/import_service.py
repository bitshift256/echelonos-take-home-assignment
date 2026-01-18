"""Import service for bulk importing data from CSV and Excel."""
import io
import csv
from typing import List, Optional
from datetime import datetime
from sqlalchemy.orm import Session
from openpyxl import load_workbook
from app.models import Employee, EmployeeStatus
from app.schemas import BulkImportResult


class ImportService:
    """Service for bulk importing employee data."""

    def __init__(self, db: Session):
        self.db = db

    def import_employees_csv(self, file_content: bytes) -> BulkImportResult:
        """Import employees from CSV file."""
        result = BulkImportResult(
            total_processed=0,
            successful=0,
            failed=0,
            errors=[]
        )
        
        try:
            content = file_content.decode('utf-8')
            reader = csv.DictReader(io.StringIO(content))
            
            for row_num, row in enumerate(reader, start=2):  # Start at 2 (after header)
                result.total_processed += 1
                try:
                    employee = self._create_employee_from_row(row)
                    self.db.add(employee)
                    result.successful += 1
                except Exception as e:
                    result.failed += 1
                    result.errors.append({
                        "row": row_num,
                        "error": str(e),
                        "data": row
                    })
            
            if result.successful > 0:
                self.db.commit()
                
        except Exception as e:
            result.errors.append({"error": f"File parsing error: {str(e)}"})
        
        return result

    def import_employees_excel(self, file_content: bytes) -> BulkImportResult:
        """Import employees from Excel file."""
        result = BulkImportResult(
            total_processed=0,
            successful=0,
            failed=0,
            errors=[]
        )
        
        try:
            wb = load_workbook(filename=io.BytesIO(file_content))
            ws = wb.active
            
            # Get headers from first row
            headers = [cell.value for cell in ws[1]]
            header_map = {h.lower().replace(" ", "_"): i for i, h in enumerate(headers) if h}
            
            for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if not any(row):  # Skip empty rows
                    continue
                    
                result.total_processed += 1
                try:
                    row_dict = {key: row[idx] for key, idx in header_map.items() if idx < len(row)}
                    employee = self._create_employee_from_row(row_dict)
                    self.db.add(employee)
                    result.successful += 1
                except Exception as e:
                    result.failed += 1
                    result.errors.append({
                        "row": row_num,
                        "error": str(e),
                        "data": dict(zip(headers, row))
                    })
            
            if result.successful > 0:
                self.db.commit()
                
        except Exception as e:
            result.errors.append({"error": f"File parsing error: {str(e)}"})
        
        return result

    def _create_employee_from_row(self, row: dict) -> Employee:
        """Create an Employee object from a row dictionary."""
        # Normalize keys (handle various naming conventions)
        normalized = {}
        key_mappings = {
            'employee_id': ['employee_id', 'emp_id', 'id', 'employee id'],
            'first_name': ['first_name', 'firstname', 'first name', 'first'],
            'last_name': ['last_name', 'lastname', 'last name', 'last'],
            'email': ['email', 'e-mail', 'email address'],
            'phone': ['phone', 'phone_number', 'phone number', 'telephone'],
            'title': ['title', 'job_title', 'job title', 'position'],
            'department': ['department', 'dept', 'department name'],
            'hire_date': ['hire_date', 'hiredate', 'hire date', 'start_date', 'start date'],
            'salary': ['salary', 'pay', 'compensation'],
            'status': ['status', 'employee_status', 'employee status'],
            'manager_id': ['manager_id', 'manager', 'reports_to', 'reports to']
        }
        
        for target_key, source_keys in key_mappings.items():
            for source_key in source_keys:
                if source_key in row and row[source_key] is not None:
                    normalized[target_key] = row[source_key]
                    break
                # Also check lowercase version of actual keys
                for actual_key in row.keys():
                    if actual_key.lower().replace(" ", "_") in source_keys or actual_key.lower() in source_keys:
                        if row[actual_key] is not None:
                            normalized[target_key] = row[actual_key]
                            break
        
        # Validate required fields
        required_fields = ['employee_id', 'first_name', 'last_name', 'email', 'title', 'department', 'hire_date']
        missing = [f for f in required_fields if f not in normalized or not normalized[f]]
        if missing:
            raise ValueError(f"Missing required fields: {', '.join(missing)}")
        
        # Parse hire_date
        hire_date = normalized['hire_date']
        if isinstance(hire_date, str):
            for fmt in ['%Y-%m-%d', '%m/%d/%Y', '%d/%m/%Y', '%Y/%m/%d']:
                try:
                    hire_date = datetime.strptime(hire_date, fmt).date()
                    break
                except ValueError:
                    continue
            else:
                raise ValueError(f"Invalid date format: {normalized['hire_date']}")
        elif hasattr(hire_date, 'date'):
            hire_date = hire_date.date()
        
        # Parse status
        status = EmployeeStatus.ACTIVE
        if 'status' in normalized:
            status_str = str(normalized['status']).lower()
            status_map = {
                'active': EmployeeStatus.ACTIVE,
                'inactive': EmployeeStatus.INACTIVE,
                'on_leave': EmployeeStatus.ON_LEAVE,
                'on leave': EmployeeStatus.ON_LEAVE,
                'terminated': EmployeeStatus.TERMINATED
            }
            status = status_map.get(status_str, EmployeeStatus.ACTIVE)
        
        # Parse salary
        salary = None
        if 'salary' in normalized and normalized['salary']:
            try:
                salary = float(str(normalized['salary']).replace(',', '').replace('$', ''))
            except ValueError:
                pass
        
        # Parse manager_id
        manager_id = None
        if 'manager_id' in normalized and normalized['manager_id']:
            try:
                manager_id = int(normalized['manager_id'])
            except (ValueError, TypeError):
                pass
        
        return Employee(
            employee_id=str(normalized['employee_id']),
            first_name=str(normalized['first_name']),
            last_name=str(normalized['last_name']),
            email=str(normalized['email']),
            phone=str(normalized.get('phone', '')) or None,
            title=str(normalized['title']),
            department=str(normalized['department']),
            hire_date=hire_date,
            salary=salary,
            status=status,
            manager_id=manager_id
        )

