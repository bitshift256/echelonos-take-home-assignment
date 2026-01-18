"""Export service for generating CSV, Excel, and PDF exports."""
import io
import csv
from typing import List, Optional
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from app.models import Employee, Team


class ExportService:
    """Service for exporting data to various formats."""

    @staticmethod
    def employees_to_csv(employees: List[Employee]) -> str:
        """Export employees to CSV format."""
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Header
        writer.writerow([
            "Employee ID", "First Name", "Last Name", "Email", "Phone",
            "Title", "Department", "Hire Date", "Salary", "Status", "Manager ID"
        ])
        
        # Data
        for emp in employees:
            writer.writerow([
                emp.employee_id, emp.first_name, emp.last_name, emp.email, emp.phone,
                emp.title, emp.department, emp.hire_date.isoformat() if emp.hire_date else "",
                emp.salary, emp.status.value if emp.status else "", emp.manager_id or ""
            ])
        
        return output.getvalue()

    @staticmethod
    def employees_to_excel(employees: List[Employee]) -> bytes:
        """Export employees to Excel format."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Employees"
        
        # Header styling
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = [
            "Employee ID", "First Name", "Last Name", "Email", "Phone",
            "Title", "Department", "Hire Date", "Salary", "Status", "Manager ID"
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # Data
        for row, emp in enumerate(employees, 2):
            data = [
                emp.employee_id, emp.first_name, emp.last_name, emp.email, emp.phone,
                emp.title, emp.department, 
                emp.hire_date.isoformat() if emp.hire_date else "",
                emp.salary, emp.status.value if emp.status else "", emp.manager_id or ""
            ]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='left')
        
        # Adjust column widths
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + col)].width = 15
        
        # Save to bytes
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def employees_to_pdf(employees: List[Employee], title: str = "Employee Directory") -> bytes:
        """Export employees to PDF format."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(letter), 
                               leftMargin=0.5*inch, rightMargin=0.5*inch)
        elements = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=20,
            alignment=1  # Center
        )
        
        # Title
        elements.append(Paragraph(title, title_style))
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", 
                                  styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Table data
        data = [[
            "ID", "Name", "Email", "Title", "Dept", "Hire Date", "Status"
        ]]
        
        for emp in employees:
            data.append([
                emp.employee_id,
                f"{emp.first_name} {emp.last_name}",
                emp.email,
                emp.title[:30] + "..." if len(emp.title) > 30 else emp.title,
                emp.department,
                emp.hire_date.strftime("%Y-%m-%d") if emp.hire_date else "",
                emp.status.value if emp.status else ""
            ])
        
        # Create table
        table = Table(data, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4E79')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
            ('BACKGROUND', (0, 1), (-1, -1), colors.white),
            ('TEXTCOLOR', (0, 1), (-1, -1), colors.black),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 1, colors.black),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F0F0')]),
        ]))
        
        elements.append(table)
        
        # Build PDF
        doc.build(elements)
        return output.getvalue()

    @staticmethod
    def teams_to_csv(teams: List[Team]) -> str:
        """Export teams to CSV format."""
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Header
        writer.writerow(["ID", "Name", "Description", "Parent Team ID", "Lead ID"])
        
        # Data
        for team in teams:
            writer.writerow([
                team.id, team.name, team.description or "",
                team.parent_team_id or "", team.lead_id or ""
            ])
        
        return output.getvalue()

    @staticmethod
    def teams_to_excel(teams: List[Team]) -> bytes:
        """Export teams to Excel format."""
        wb = Workbook()
        ws = wb.active
        ws.title = "Teams"
        
        # Header styling
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ["ID", "Name", "Description", "Parent Team ID", "Lead ID"]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border
        
        # Data
        for row, team in enumerate(teams, 2):
            data = [team.id, team.name, team.description or "", 
                   team.parent_team_id or "", team.lead_id or ""]
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.border = thin_border
        
        # Adjust column widths
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 50
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 10
        
        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    @staticmethod
    def org_chart_to_pdf(employees: List[Employee], title: str = "Organization Chart") -> bytes:
        """Export org chart to PDF format."""
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=letter)
        elements = []
        
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=20,
            alignment=1
        )
        
        elements.append(Paragraph(title, title_style))
        elements.append(Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", 
                                  styles['Normal']))
        elements.append(Spacer(1, 20))
        
        # Build hierarchy
        def build_hierarchy(manager_id: Optional[int], level: int = 0) -> List:
            """Recursively build org chart hierarchy."""
            items = []
            direct_reports = [e for e in employees if e.manager_id == manager_id]
            
            for emp in direct_reports:
                indent = "    " * level
                name = f"{emp.first_name} {emp.last_name}"
                items.append(Paragraph(
                    f"{indent}• <b>{name}</b> - {emp.title} ({emp.department})",
                    styles['Normal']
                ))
                items.extend(build_hierarchy(emp.id, level + 1))
            
            return items
        
        # Start with employees who have no manager
        root_employees = [e for e in employees if e.manager_id is None]
        for emp in root_employees:
            elements.append(Paragraph(
                f"<b>{emp.first_name} {emp.last_name}</b> - {emp.title} ({emp.department})",
                styles['Heading2']
            ))
            elements.extend(build_hierarchy(emp.id, 1))
            elements.append(Spacer(1, 10))
        
        doc.build(elements)
        return output.getvalue()

